"""Техническое приложение и доклад из Markdown в Word.

Оформление как у резюме и заявки (resume.py): Times New Roman, A4, таблицы с сеткой и серой шапкой.
Понимает заголовки #, ##, ###, абзацы, списки «- » и «1. », таблицы и **жирный**, *курсив*, `код`.

    .venv\\Scripts\\python -m docs.submission.to_docx
"""

from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from docs.submission.resume import FONT, HERE, _fixed_widths, _font, _para

# файл → (результат, размер основного шрифта); one-pager собирается отдельно: one_pager.py
FILES = {"tech_appendix_ru.md": ("tech_appendix_ru.docx", 11), "doklad_ru.md": ("doklad_ru.docx", 12)}
TEXT_WIDTH_CM = 17.5
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\s][^*]*?\*(?!\*))")
HEADING_SIZE = {1: 15, 2: 13, 3: 11.5}


def inline(par, text: str, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    """**жирный**, *курсив*, `код` внутри абзаца."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            inline(par, part[2:-2], size, bold=True, italic=italic)
        elif part.startswith("`"):
            run = _font(par.add_run(part[1:-1]), size - 1, bold=bold, italic=italic)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif part.startswith("*"):
            inline(par, part[1:-1], size, bold=bold, italic=True)
        else:
            _font(par.add_run(part), size, bold=bold, italic=italic)


def _shade(cell, color: str = "E7E9EC") -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _widths(rows: list[list[str]]) -> list:
    """Каждой колонке: минимум, чтобы не рвалось самое длинное слово; остаток ширины: по объёму текста."""
    char_cm, pad_cm = 0.2, 0.35  # ширина символа 9,5 pt жирным с запасом и поля ячейки
    minimum, volume = [], []
    for j in range(len(rows[0])):
        cells = [re.sub(r"[*`]", "", r[j]) for r in rows if j < len(r)]
        longest_word = max((len(w) for c in cells for w in c.split()), default=1)
        minimum.append(max(longest_word, 3) * char_cm + pad_cm)
        volume.append(sum(min(len(c), 60) for c in cells) / len(cells))
    spare = max(TEXT_WIDTH_CM - sum(minimum), 0)
    return [Cm(m + spare * v / sum(volume)) for m, v in zip(minimum, volume)]


def _repeat_header(row) -> None:
    """Шапка таблицы повторяется на каждой странице."""
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(flag)


def table(doc, lines: list[str]) -> None:
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    grid = doc.add_table(rows=len(rows), cols=len(rows[0]), style="Table Grid")
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_widths(grid, _widths(rows))
    _repeat_header(grid.rows[0])
    for i, row in enumerate(rows):
        for j, text in enumerate(row[:len(rows[0])]):
            cell = grid.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            inline(par, text.replace(" ± ", " ± "), size=9.5, bold=i == 0)  # «0,315 ± 0,003» не рвётся
            if i == 0:
                _shade(cell)
    _para(doc, after=2)


def build(md: str, out, size: float = 11) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    section.left_margin, section.right_margin = Cm(2.0), Cm(1.5)
    doc.styles["Normal"].font.name = FONT

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        heading = re.match(r"(#{1,3}) (.+)", line)
        if heading:
            level = len(heading.group(1))
            par = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT,
                        before=0 if level == 1 else 8, after=4)
            par.paragraph_format.keep_with_next = True
            inline(par, heading.group(2), size=HEADING_SIZE[level], bold=True)
            i += 1
            continue
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(doc, block)
            continue
        item = re.match(r"\s*(- |\d+\. )(.+)", line)
        if item:
            marker = "•" if item.group(1) == "- " else item.group(1).strip()
            par = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
            fmt = par.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Cm(0.6), Cm(-0.6)
            fmt.tab_stops.add_tab_stop(Cm(0.6))
            inline(par, f"{marker}\t{item.group(2)}", size)
            i += 1
            continue
        text = [line.strip()]  # абзац: строки до пустой строки, списка, таблицы или заголовка
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"\s*(#{1,3} |\||- |\d+\. )", lines[i]):
            text.append(lines[i].strip())
            i += 1
        if any("http" in line for line in text):  # строки со ссылками: каждая отдельно, без растяжки
            par = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=5)
            for k, line in enumerate(text):
                if k:
                    par.add_run().add_break()
                inline(par, line, size)
        else:
            inline(_para(doc, after=5), " ".join(text), size)
    doc.save(out)


def main() -> int:
    for src, (dst, size) in FILES.items():
        out = HERE / dst
        build((HERE / src).read_text(encoding="utf-8"), out, size)
        print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

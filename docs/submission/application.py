"""Заявка от ОВПО на участие в конкурсе (п. 8.1 положения): документ Word для подписи и печати.

Личные данные участника (телефон, e-mail, курс) берутся из application_private.json: этот файл,
как и готовый application_ru.docx, в git не попадает. Поля, которые заполняет вуз, подсвечены жёлтым.

    .venv\\Scripts\\python -m docs.submission.application
"""

from __future__ import annotations

import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

from docs.submission.resume import FONT, HERE, REPO, STAND, _fixed_widths, _font, _no_borders, _para

FILL = "[заполнить]"
UNIVERSITY = "НАО «Западно-Казахстанский аграрно-технический университет имени Жангир хана»"
ADDRESS = "090009, Республика Казахстан, Западно-Казахстанская область, г. Уральск, ул. Жангир хана, 51"
PROJECT = "ШЕКАРА: межбанковское выявление сетей дропов без обмена клиентскими данными"


def _runs(par, *parts, size=11):
    """parts: строки или (текст, "fill") для подсвеченного поля, (текст, "bold") для жирного."""
    for part in parts:
        text, kind = part if isinstance(part, tuple) else (part, "")
        run = _font(par.add_run(text), size, bold=kind == "bold")
        if kind == "fill":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def _table(doc, rows, widths=(Cm(5.2), Cm(12.3))):
    table = doc.add_table(rows=0, cols=2, style="Table Grid")
    for label, value in rows:
        cells = table.add_row().cells
        _runs(cells[0].paragraphs[0], (label, "bold"), size=10.5)
        values = value if isinstance(value, list) else [value]
        _runs(cells[1].paragraphs[0], *values, size=10.5)
        for cell in cells:
            fmt = cell.paragraphs[0].paragraph_format
            fmt.space_before, fmt.space_after = Pt(1), Pt(1)
    _fixed_widths(table, widths)
    return table


def build(person: dict):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    section.left_margin, section.right_margin = Cm(2.0), Cm(1.5)
    doc.styles["Normal"].font.name = FONT

    # адресат справа
    for line in ("В Организационный комитет Республиканского конкурса",
                 "ОЮЛ «Ассоциация банков Республики Казахстан»",
                 "г. Алматы, ул. Панфилова, 98, блок «А», офис 518",
                 "info@abrk.kz"):
        par = _para(doc, line, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
        par.paragraph_format.left_indent = Cm(8.0)

    _para(doc, "ЗАЯВКА", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=0)
    _para(doc, "на участие в Республиканском конкурсе ОЮЛ «Ассоциация банков Республики Казахстан» 2026",
          size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    _runs(_para(doc, after=8),
          f"{UNIVERSITY} выдвигает для участия в конкурсе работу «{PROJECT}» "
          "(треки AI for Finance (антифрод) и Digital Banking (AML/regtech)).")

    _para(doc, "1. Организация высшего и послевузовского образования", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
    _table(doc, [
        ("Полное наименование", UNIVERSITY),
        ("Полный адрес", ADDRESS),
        ("Телефон", (FILL, "fill")),
        ("E-mail", (FILL, "fill")),
        ("Сайт", "wkatu.edu.kz"),
    ])

    _para(doc, "2. Участник", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3)
    _table(doc, [
        ("Ф. И. О.", "Джусупов Сагир Едгеевич"),
        ("Статус", person["status"]),
        ("Институт", person["institute"]),
        ("Образовательная программа", person["program"]),
        ("Роль в работе", "автор: протокол, модель, эксперимент, веб-стенд, документация"),
        ("Телефон", person["phone"]),
        ("E-mail", person["email"]),
    ])

    _para(doc, "3. Работа", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3)
    _table(doc, [
        ("Название", f"«{PROJECT}»"),
        ("Тип работы", "MVP с демо и аналитическое решение с экспериментальными метриками"),
        ("Научный руководитель", ["Касымова Акмарал Хамзиевна, ", ("[должность, учёная степень]", "fill")]),
        ("Работающий прототип", STAND),
        ("Исходный код", REPO),
    ])

    _para(doc, "4. Прилагаемые материалы", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3)
    for i, item in enumerate((
            "Резюме проекта на русском и английском языках (по одной странице).",
            "One-pager: проблема, целевая аудитория, решение, ценность, новизна.",
            "Презентация (pitch-deck, 12 слайдов) и доклад.",
            "Техническое приложение: методика и модель, данные, метрики, план доработки (TRL).",
            "Демо: работающий прототип в интернете (ссылка выше)."), 1):
        par = _para(doc, f"{i}. {item}", align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
        par.paragraph_format.left_indent = Cm(0.5)

    signs = doc.add_table(rows=0, cols=2)
    _no_borders(signs)
    for role, name in ((("[Должность руководителя ОВПО]", "fill"), ("[Ф. И. О.]", "fill")),
                       ("Научный руководитель", "Касымова А. Х.")):
        cells = signs.add_row().cells
        _runs(cells[0].paragraphs[0], role)
        _runs(cells[1].paragraphs[0], "____________________  / ", name, " /")
    cells = signs.add_row().cells
    _runs(cells[0].paragraphs[0], "М. П.")
    _runs(cells[1].paragraphs[0], "«____» _____________ 2026 г.")
    for row in signs.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_before = Pt(10)
    _fixed_widths(signs, (Cm(6.5), Cm(11.0)))

    out = HERE / "application_ru.docx"
    doc.save(out)
    return out


def main() -> int:
    person = json.loads((HERE / "application_private.json").read_text(encoding="utf-8"))
    print(build(person))
    return 0


if __name__ == "__main__":
    sys.exit(main())

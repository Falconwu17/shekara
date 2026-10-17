"""One-pager для заявки (п. 8 положения): проблема, целевая аудитория, решение, ценность, новизна.

Одна страница A4 в цветах презентации: тёмная шапка, полоса ключевых цифр, пять разделов, подвал со ссылками.

    .venv\\Scripts\\python -m docs.submission.one_pager
"""

from __future__ import annotations

import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docs.submission.resume import HERE, REPO, STAND, _fixed_widths, _no_borders

FONT = "Arial"
NAVY, ACCENT, CREAM = "14213D", "A84A14", "F6F4EE"
INK, MUTED, LIGHT = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x3D, 0x4A, 0x5C), RGBColor(0xF6, 0xF4, 0xEE)
ORANGE, SOFT = RGBColor(0xA8, 0x4A, 0x14), RGBColor(0xBF, 0xC8, 0xD9)
WIDTH_CM = 18.0

NUMBERS = [
    ("≈ 45 млрд ₸", "ущерб от интернет-мошенничества в РК за 2024 год"),
    ("28%", "схемы отмывания видит в среднем один банк"),
    ("×2", "точнее центрального оператора: PR-AUC 0,315 против 0,162"),
    ("0,7 мс", "проверка входящего перевода"),
]

SECTIONS = [
    ("Проблема", [
        "Деньги жертвы социальной инженерии за минуты проходят через 3-5 счетов-дропов в разных банках. "
        "Каждый банк видит одно звено и считает его обычным переводом. В 2024 году в Казахстане более 17 тыс. "
        "случаев интернет-мошенничества, вернуть удалось около 2 из 45 млрд ₸ (МВД). Объединить данные банков "
        "не позволяют банковская тайна и закон о персональных данных."]),
    ("Целевая аудитория", [
        "Антифрод- и AML-подразделения банков второго уровня; Антифрод-центр НБК как оператор межбанковского "
        "обмена; АРРФР как владелец правил раскрытия; далее МФО через ГКБ."]),
    ("Решение", [
        "Протокол и графовая модель. Банк сообщает соседу не данные клиента, а одну оценку риска счёта и только "
        "по ссылке на реальный зачисленный перевод между их клиентами. Банк-получатель проверяет зачисление за "
        "0,7 мс, подозрительная цепочка трассируется дальше по переводам. Полную цепочку видит только оператор "
        "с мандатом. Каждое обращение пишется в журнал с подписью Ed25519, уровень раскрытия выбирает регулятор."]),
    ("Ценность", [
        "Измерено на IBM AMLworld: 5,08 млн переводов, 30 банков, 3 запуска.",
        "• Модель на графе самого банка вдвое сильнее центрального оператора, значимо в 3 из 3 запусков.",
        "• Обмен оценками добавляет ещё 13% PR-AUC во всех запусках; значимость проверит пилот.",
        "• С задержкой исходящих для верхушки очереди деньги остаются на счёте в 30% звеньев, "
        "ценой меньше одного честного клиента на банк в сутки.",
        "• Клиентские данные не покидают банк. MVP стоит 0 ₸, пилот на два банка около 30-45 млн ₸."]),
    ("Новизна", [
        "Мир решает задачу тремя путями: центральный граф у оператора платежей (Mastercard MITS), обмен "
        "записями о клиентах (COSMIC, Сингапур), общая база (TMNL, Нидерланды). ШЕКАРА предлагает для "
        "Казахстана другой путь: оценка риска только по ссылке на реальный перевод, данные остаются в банке. "
        "Подход сравнён с MITS на одном бенчмарке, цена раскрытия измерена в качестве и в ε."]),
]


def _run(par, text, size=10, bold=False, color=INK):
    run = par.add_run(text)
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size, run.bold, run.font.color.rgb = Pt(size), bold, color
    return run


def _par(cell, text="", *, size=10, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, first=False):
    par = cell.paragraphs[0] if first else cell.add_paragraph()
    par.alignment = align
    fmt = par.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(after), 1.1
    if text:
        _run(par, text, size, bold, color)
    return par


def _fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _margins(cell, top=0.25, side=0.35):
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", top), ("start", side), ("end", side)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(value * 567)))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    cell._tc.get_or_add_tcPr().append(margins)


def _left_rule(cell, color=ACCENT):
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:start")
    for key, value in (("w:val", "single"), ("w:sz", "18"), ("w:color", color)):
        left.set(qn(key), value)
    borders.append(left)
    cell._tc.get_or_add_tcPr().append(borders)


def _block(doc, cols):
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(table)
    return table


def _gap(doc, pt):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = Pt(pt)
    return par


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.2), Cm(1.0)
    section.left_margin = section.right_margin = Cm(1.5)

    # шапка
    head = _block(doc, 1)
    cell = head.cell(0, 0)
    _fill(cell, NAVY)
    _margins(cell, top=0.45, side=0.6)
    _par(cell, "КОНКУРС АБРК 2026 · AI FOR FINANCE · DIGITAL BANKING", size=8.5, bold=True, color=RGBColor(0xF0, 0xA3, 0x5E), first=True, after=4)
    _par(cell, "ШЕКАРА", size=30, bold=True, color=LIGHT, after=2)
    _par(cell, "Межбанковское выявление сетей дропов без обмена клиентскими данными", size=13, color=LIGHT, after=0)
    _fixed_widths(head, [Cm(WIDTH_CM)])
    _gap(doc, 11)

    # ключевые цифры
    kpi = _block(doc, len(NUMBERS))
    for cell, (value, label) in zip(kpi.rows[0].cells, NUMBERS):
        _fill(cell, CREAM)
        _margins(cell, top=0.2, side=0.3)
        _left_rule(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _par(cell, value, size=19, bold=True, color=ORANGE, first=True, after=1)
        _par(cell, label, size=9, color=MUTED)
    _fixed_widths(kpi, [Cm(WIDTH_CM / len(NUMBERS))] * len(NUMBERS))
    _gap(doc, 11)

    # пять разделов из положения
    body = doc.add_table(rows=0, cols=2)
    body.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(body)
    for title, lines in SECTIONS:
        label, text = body.add_row().cells
        _margins(label, top=0.24, side=0.2)
        _margins(text, top=0.24, side=0.3)
        _left_rule(label)
        _par(label, title.upper(), size=9, bold=True, color=ORANGE, first=True)
        for i, line in enumerate(lines):
            par = _par(text, line, size=10.5, color=INK if i == 0 else MUTED, first=i == 0, after=1)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if line.startswith("•"):
                par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.35), Cm(-0.35)
    _fixed_widths(body, (Cm(3.6), Cm(WIDTH_CM - 3.6)))
    _gap(doc, 11)

    # подвал
    foot = _block(doc, 1)
    cell = foot.cell(0, 0)
    _fill(cell, NAVY)
    _margins(cell, top=0.3, side=0.6)
    par = _par(cell, first=True, after=2)
    _run(par, "Джусупов Сагир Едгеевич", 9.5, True, LIGHT)
    _run(par, ", магистрант ЗКАТУ имени Жангир хана, Уральск", 9.5, False, SOFT)
    par = _par(cell, after=2)
    _run(par, "Научный руководитель: ", 9.5, False, SOFT)
    _run(par, "Касымова Акмарал Хамзиевна", 9.5, True, LIGHT)
    par = _par(cell)
    _run(par, "Работающий стенд: ", 9.5, False, SOFT)
    _run(par, STAND.removeprefix("https://"), 9.5, True, LIGHT)
    _run(par, "   ·   Открытый код: ", 9.5, False, SOFT)
    _run(par, REPO.removeprefix("https://"), 9.5, True, LIGHT)
    _fixed_widths(foot, [Cm(WIDTH_CM)])

    out = HERE / "one_pager_ru.docx"
    doc.save(out)
    return out


def main() -> int:
    print(build())
    return 0


if __name__ == "__main__":
    sys.exit(main())
